"""LLM 판사 타당성 검증용 파일 생성.

출력:
  mafs/evaluate/results/validation/
    reading_samples.docx   — 채점자용 읽기 파일 (에세이 + 피드백 20개)
    scoring_input.xlsx     — 채점 입력 파일 (LLM 점수 + R1·R2 빈 칸)
"""

import json
import random
from pathlib import Path
from collections import defaultdict

# ── 상수 ──────────────────────────────────────────────────────────────────────
RQ1_PATH   = "mafs/evaluate/results/rq1_200.jsonl"
OUT_DIR    = Path("mafs/evaluate/results/validation")
SEED       = 42
N_PER_CELL = 2  # 학년×교과 조합당 샘플 수

CRITERIA   = ["C1", "C2", "C3", "C4", "C5", "M1", "M2", "M3", "M4", "M5"]
CRIT_NAMES = {
    "C1": "글쓰기 관련성", "C2": "약점 구체성", "C3": "강점 근거",
    "C4": "분석 근거 수치화", "C5": "학습 성장 전제",
    "M1": "피드포워드", "M2": "분석 타당성",
    "M3": "학습자 수준 적합", "M4": "과제 맥락 반영", "M5": "영역별 분석의 충실성",
}


# ── 1. 샘플 추출 ───────────────────────────────────────────────────────────────
def load_and_sample() -> list[dict]:
    records = []
    with open(RQ1_PATH, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                records.append(json.loads(line))

    # 학년×교과 그룹화
    groups: dict[tuple, list] = defaultdict(list)
    for r in records:
        key = (r.get("grade", ""), r.get("subject", ""))
        groups[key].append(r)

    rng = random.Random(SEED)
    sampled = []
    for key in sorted(groups.keys()):
        pool = groups[key]
        n = min(N_PER_CELL, len(pool))
        sampled.extend(rng.sample(pool, n))

    # 번호 부여
    for i, rec in enumerate(sampled, 1):
        rec["_sample_no"] = i
    return sampled


# ── 2. Word 읽기용 파일 ────────────────────────────────────────────────────────
def make_word(samples: list[dict], out_path: Path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin    = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin   = Cm(3)
    sec.right_margin  = Cm(3)

    # 문서 제목
    h = doc.add_paragraph("LLM 판사 타당성 검증 — 채점용 샘플 자료")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.runs[0].bold = True
    h.runs[0].font.size = Pt(14)

    notice = doc.add_paragraph(
        "※ 각 샘플의 과제 지시문과 학생 에세이를 먼저 읽은 뒤, "
        "MAFS 피드백을 채점 기준에 따라 평가해 주세요.\n"
        "채점 대상은 학생 글의 수준이 아닌 피드백의 품질입니다."
    )
    notice.runs[0].font.size = Pt(9)
    doc.add_paragraph()

    for rec in samples:
        no       = rec["_sample_no"]
        grade    = rec.get("grade", "")
        subject  = rec.get("subject", "")
        purpose  = rec.get("purpose", "")
        prompt   = rec.get("prompt", "").strip()
        essay    = rec.get("essay", "").strip()
        feedback = rec["condition_a"].get("feedback", "").strip()

        # 샘플 헤더
        hdr = doc.add_paragraph(f"샘플 {no:02d}   |   {grade} · {subject}   |   글쓰기 목적: {purpose}")
        hdr.runs[0].bold = True
        hdr.runs[0].font.size = Pt(11)
        hdr.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

        doc.add_paragraph("─" * 55).runs[0].font.size = Pt(8)

        # 과제 지시문
        _section_title(doc, "과제 지시문")
        p = doc.add_paragraph(prompt)
        p.runs[0].font.size = Pt(9)
        doc.add_paragraph()

        # 학생 에세이
        _section_title(doc, "학생 에세이")
        p = doc.add_paragraph(essay)
        p.runs[0].font.size = Pt(9)
        doc.add_paragraph()

        # MAFS 피드백
        _section_title(doc, "MAFS 피드백  ← 채점 대상")
        p = doc.add_paragraph(feedback)
        p.runs[0].font.size = Pt(9)

        # 구분선
        doc.add_paragraph()
        doc.add_paragraph("━" * 55).runs[0].font.size = Pt(8)
        doc.add_page_break()

    doc.save(out_path)
    print(f"[Word] {out_path}")


def _section_title(doc, title: str):
    from docx.shared import Pt, RGBColor
    p = doc.add_paragraph(f"[ {title} ]")
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.color.rgb = RGBColor(0x70, 0x30, 0xA0)


# ── 3. Excel 채점 입력 파일 ───────────────────────────────────────────────────
def make_excel(samples: list[dict], out_path: Path):
    import openpyxl
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side
    )

    wb = openpyxl.Workbook()

    # ── 시트 1: 채점 입력 ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "채점입력"

    # 색 정의
    fill_header  = PatternFill("solid", fgColor="2E4057")
    fill_meta    = PatternFill("solid", fgColor="D9E1F2")
    fill_llm     = PatternFill("solid", fgColor="E2EFDA")
    fill_r1      = PatternFill("solid", fgColor="FFF2CC")
    fill_r2      = PatternFill("solid", fgColor="FCE4D6")
    fill_total   = PatternFill("solid", fgColor="F2F2F2")
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    def hdr_font(color="FFFFFF", bold=True, size=9):
        return Font(color=color, bold=bold, size=size)

    def cell_font(bold=False, size=9):
        return Font(bold=bold, size=size)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    # ── 헤더 행 1: 그룹 레이블 ──────────────────────────────────────────────
    groups = [
        ("A1:D1", "샘플 정보",     fill_meta),
        ("E1:N1", "LLM 판사 점수 (참고용)", fill_llm),
        ("O1:X1", "평가자 1 (R1) 점수",     fill_r1),
        ("Y1:AH1", "평가자 2 (R2) 점수",    fill_r2),
        ("AI1:AL1", "합계",                  fill_total),
    ]
    for rng, label, fill in groups:
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value = label
        c.fill  = fill
        c.font  = hdr_font(color="1F497D" if fill != fill_header else "FFFFFF", size=9)
        c.alignment = center
        c.border = thin

    # ── 헤더 행 2: 컬럼명 ───────────────────────────────────────────────────
    row2_labels = (
        ["샘플", "학년", "교과", "글쓰기 목적"]
        + [f"LLM\n{c}" for c in CRITERIA]
        + [f"R1\n{c}" for c in CRITERIA]
        + [f"R2\n{c}" for c in CRITERIA]
        + ["LLM\n합계", "R1\n합계", "R2\n합계", "비고"]
    )
    fills_row2 = (
        [fill_meta] * 4
        + [fill_llm] * 10
        + [fill_r1]  * 10
        + [fill_r2]  * 10
        + [fill_total] * 4
    )
    for col_i, (label, fill) in enumerate(zip(row2_labels, fills_row2), 1):
        c = ws.cell(row=2, column=col_i, value=label)
        c.fill      = fill
        c.font      = hdr_font(color="000000", size=8)
        c.alignment = center
        c.border    = thin

    # ── 데이터 행 ────────────────────────────────────────────────────────────
    for rec in samples:
        row = rec["_sample_no"] + 2  # 헤더 2행 아래
        scores = rec["condition_a"].get("scores", {})

        meta = [
            rec["_sample_no"],
            rec.get("grade", ""),
            rec.get("subject", ""),
            rec.get("purpose", ""),
        ]
        llm_scores = [scores.get(c, "") for c in CRITERIA]
        llm_total  = scores.get("score_total", "")

        row_data = meta + llm_scores + [""] * 10 + [""] * 10 + [llm_total, "", "", ""]
        row_fills = (
            [fill_meta] * 4
            + [fill_llm] * 10
            + [fill_r1]  * 10
            + [fill_r2]  * 10
            + [fill_total] * 4
        )

        for col_i, (val, fill) in enumerate(zip(row_data, row_fills), 1):
            c = ws.cell(row=row, column=col_i, value=val)
            c.fill      = fill
            c.font      = cell_font(size=9)
            c.alignment = center if col_i != 4 else left
            c.border    = thin

        # R1·R2 합계 수식
        r1_start = openpyxl.utils.get_column_letter(5 + 10)     # O
        r1_end   = openpyxl.utils.get_column_letter(5 + 10 + 9) # X
        r2_start = openpyxl.utils.get_column_letter(5 + 20)     # Y
        r2_end   = openpyxl.utils.get_column_letter(5 + 20 + 9) # AH
        ws.cell(row=row, column=36).value = f"=SUM({r1_start}{row}:{r1_end}{row})"
        ws.cell(row=row, column=37).value = f"=SUM({r2_start}{row}:{r2_end}{row})"

    # 열 너비
    col_widths = [6, 6, 6, 10] + [6] * 30 + [7, 7, 7, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 28
    for i in range(3, len(samples) + 3):
        ws.row_dimensions[i].height = 18

    ws.freeze_panes = "E3"

    # ── 시트 2: 기준 설명 ─────────────────────────────────────────────────────
    ws2 = wb.create_sheet("기준 설명")
    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 35
    ws2.column_dimensions["D"].width = 35
    ws2.column_dimensions["E"].width = 35

    ref_headers = ["기준", "명칭", "0점 (미충족)", "1점 (부분 충족)", "2점 (완전 충족)"]
    ref_data = [
        ("C1", "글쓰기 관련성",
         "에세이 표현을 전혀 인용하지 않거나 1개 이하",
         "에세이 표현 참조는 있으나 따옴표 직접 인용 2개 이하",
         "에세이 표현·문장 5개↑, 그 중 3개↑ 따옴표 인용, 각각 다른 피드백 지점에 연결"),
        ("C2", "약점 구체성",
         "약점 지적 없거나 '부족하다' 수준의 막연한 서술",
         "'위치+이유+인용' 중 일부만 갖추거나, 조건 갖추었으나 1~2건만",
         "'위치+이유+에세이 해당 표현 직접 인용' 모두 갖춘 구체적 약점 3건↑"),
        ("C3", "강점 근거",
         "강점 언급 없거나 '잘 썼다' 수준의 막연한 칭찬",
         "실제 표현 인용 1개이거나, 인용만 있고 '왜 강점인지' 설명 없음",
         "에세이 실제 표현·구조 2개↑ 직접 인용 + 각 인용에 '왜 강점인지' 설명"),
        ("C4", "분석 근거 수치화",
         "도구 분석 수치(맞춤법 오류 건수, 접속어 유형 수 등)가 전혀 없음",
         "수치가 있으나 학년 기준 없거나 해석 없이 나열만",
         "도구 분석 수치 1개↑ 명시 + 학년 기준과 함께 구체적 제시"),
        ("C5", "학습 성장 전제",
         "약점 나열만 있고 강점·피드포워드 없음",
         "강점·개선점·피드포워드 중 일부만 있어 균형 무너짐",
         "강점·개선점·피드포워드가 균형 있게 구성, 학습 의욕 충분히 지지"),
        ("M1", "피드포워드",
         "'다음에는 ~해보세요' 형식의 구체적 지침 없음",
         "개선 방향은 있으나 에세이 특정 약점과 연결 안 되거나 형식 모호",
         "에세이의 특정 약점과 연결된 다음 글쓰기 지침이 '다음에는 ~해보세요' 형식으로 명시"),
        ("M2", "분석 타당성",
         "에세이에 없는 내용 언급 또는 명백한 오분석 1건↑",
         "타당한 분석이 대부분이나 경미한 과장·왜곡 1건",
         "지적한 약점·강점이 에세이 실제 내용에 비추어 모두 타당"),
        ("M3", "학습자 수준 적합",
         "해당 학년이 이해하기 어려운 어휘·문체·어조",
         "대체로 적합하나 일부 어휘·표현이 학년 수준과 맞지 않음",
         "어휘·문장·어조 전반이 해당 학년 학생이 충분히 이해·수용 가능"),
        ("M4", "과제 맥락 반영",
         "과제 지시문 조건을 전혀 언급하지 않거나 단순 나열",
         "지시문 조건 1가지만 에세이 내용과 연결하여 분석",
         "과제 지시문 조건을 에세이 내용과 연결하여 2가지↑ 분석 (단순 나열 불인정)"),
        ("M5", "영역별 분석의 충실성",
         "4개 영역(과제수행·내용·조직·표현) 중 2개↑ 누락",
         "4개 영역이 있으나 일부 영역에 에세이 근거 포함 분석이 1문장 이하",
         "4개 영역 각각에서 에세이 근거를 포함한 분석이 2문장↑"),
    ]

    for ci, h in enumerate(ref_headers, 1):
        c = ws2.cell(row=1, column=ci, value=h)
        c.fill = PatternFill("solid", fgColor="2E4057")
        c.font = Font(color="FFFFFF", bold=True, size=9)
        c.alignment = center
        c.border = thin

    for ri, row_data in enumerate(ref_data, 2):
        bg = "DDEEFF" if row_data[0].startswith("C") else "FFE8CC"
        for ci, val in enumerate(row_data, 1):
            c = ws2.cell(row=ri, column=ci, value=val)
            c.fill      = PatternFill("solid", fgColor=bg)
            c.font      = Font(size=9)
            c.alignment = Alignment(horizontal="center" if ci == 1 else "left",
                                    vertical="center", wrap_text=True)
            c.border = thin
        ws2.row_dimensions[ri].height = 40

    wb.save(out_path)
    print(f"[Excel] {out_path}")


# ── main ──────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    samples = load_and_sample()
    print(f"샘플 {len(samples)}건 추출 완료")
    for s in samples:
        print(f"  {s['_sample_no']:2d}. {s['grade']} {s['subject']} / {s['file']}")

    make_word(samples, OUT_DIR / "reading_samples.docx")
    make_excel(samples, OUT_DIR / "scoring_input.xlsx")
    print("완료.")
