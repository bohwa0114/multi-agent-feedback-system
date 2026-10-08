"""LLM 판사 타당성 검증 결과 정리 Excel 생성.

시트 구성:
  R1_입력   — 평가자 1 점수 입력 (1~5점)
  R2_입력   — 평가자 2 점수 입력 (1~5점)
  결과요약  — M·SD·I-CVI·S-CVI·상관 자동 계산

출력: mafs/evaluate/results/validation/결과정리.xlsx
"""

import json
import random
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RQ1_PATH = "mafs/evaluate/results/rq1_200.jsonl"
OUT_PATH = Path("mafs/evaluate/results/validation/결과정리.xlsx")
SEED     = 42

CRITERIA = ["C1","C2","C3","C4","C5","M1","M2","M3","M4","M5"]
CRIT_NAMES = {
    "C1":"글쓰기 관련성", "C2":"약점 구체성",    "C3":"강점 근거",
    "C4":"분석 근거 수치화","C5":"학습 성장 전제",
    "M1":"피드포워드", "M2":"분석 타당성",
    "M3":"학습자 수준 적합", "M4":"과제 맥락 반영", "M5":"영역별 분석의 충실성",
}

# ── 스타일 ─────────────────────────────────────────────────────────────────────
def fill(hex_): return PatternFill("solid", fgColor=hex_)
def font(bold=False, size=9, color="000000"):
    return Font(bold=bold, size=size, color=color)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
thin   = Border(*[Side(style="thin")]*0,
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"),  bottom=Side(style="thin"))

F_HDR   = fill("2E4057"); F_META  = fill("D9E1F2")
F_LLM   = fill("E2EFDA"); F_R1    = fill("FFF2CC")
F_R2    = fill("FCE4D6"); F_CALC  = fill("F2F2F2")
F_WHITE = fill("FFFFFF"); F_CVI_OK= fill("C6EFCE"); F_CVI_NG= fill("FFCCCC")


def hdr_cell(ws, row, col, val, bg=None, bold=True, size=9,
             color="FFFFFF", align=center):
    c = ws.cell(row=row, column=col, value=val)
    if bg: c.fill = bg
    c.font      = font(bold=bold, size=size, color=color)
    c.alignment = align
    c.border    = thin
    return c


def data_cell(ws, row, col, val=None, bg=None,
              bold=False, size=9, align=center, formula=None):
    c = ws.cell(row=row, column=col, value=val if formula is None else formula)
    if bg: c.fill = bg
    c.font      = font(bold=bold, size=size)
    c.alignment = align
    c.border    = thin
    return c


# ── 샘플 추출 ─────────────────────────────────────────────────────────────────
def load_samples():
    records = []
    with open(RQ1_PATH, encoding="utf-8") as f:
        for line in f:
            if line.strip():
                records.append(json.loads(line.strip()))
    groups = defaultdict(list)
    for r in records:
        groups[(r.get("grade",""), r.get("subject",""))].append(r)
    rng = random.Random(SEED)
    sampled = []
    for key in sorted(groups.keys()):
        sampled.extend(rng.sample(groups[key], min(1, len(groups[key]))))
    for i, r in enumerate(sampled, 1):
        r["_no"] = i
    return sampled


# ── 입력 시트 공통 생성 ────────────────────────────────────────────────────────
def make_input_sheet(ws, samples, rater_label, bg_score):
    ws.sheet_view.showGridLines = False

    # 열 너비
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 7
    ws.column_dimensions["C"].width = 7
    ws.column_dimensions["D"].width = 10
    for i in range(5, 5 + len(CRITERIA)):
        ws.column_dimensions[get_column_letter(i)].width = 9
    ws.column_dimensions[get_column_letter(5 + len(CRITERIA))].width = 10

    # ── 행 1: 안내 ──────────────────────────────────────────────────────────
    ws.merge_cells(f"A1:{get_column_letter(5+len(CRITERIA))}1")
    c = ws["A1"]
    c.value = (f"LLM 판사 채점 적절성 검토 — {rater_label}  "
               "│  척도: 1=전혀 부적절  2=부적절  3=보통  4=적절  5=매우 적절")
    c.fill      = fill("1F3864")
    c.font      = font(bold=True, size=10, color="FFFFFF")
    c.alignment = center
    ws.row_dimensions[1].height = 22

    # ── 행 2: 그룹 헤더 ─────────────────────────────────────────────────────
    for ci, (label, bg, span) in enumerate([
        ("샘플 정보", F_META, 4),
        ("LLM 판사 점수 (참고)",    F_LLM,   len(CRITERIA)),
        (f"{rater_label} 적절성 점수 (1~5점 입력)", bg_score, len(CRITERIA)),
    ]):
        start = 1 + sum(s for _, _, s in [
            ("샘플 정보", F_META, 4),
            ("LLM 판사 점수 (참고)", F_LLM, len(CRITERIA)),
            (f"{rater_label} 적절성 점수", bg_score, len(CRITERIA)),
        ][:ci])
        end = start + span - 1
        ws.merge_cells(
            start_row=2, start_column=start,
            end_row=2,   end_column=end
        )
        c = ws.cell(row=2, column=start, value=label)
        c.fill = bg; c.font = font(bold=True, size=9, color="1F3864")
        c.alignment = center; c.border = thin
    ws.row_dimensions[2].height = 18

    # ── 행 3: 컬럼 헤더 ─────────────────────────────────────────────────────
    meta_labels = ["No.", "학년", "교과", "글쓰기\n목적"]
    for ci, lbl in enumerate(meta_labels, 1):
        hdr_cell(ws, 3, ci, lbl, F_META, color="1F3864")
    for ci, crit in enumerate(CRITERIA, 5):
        hdr_cell(ws, 3, ci,      f"LLM\n{crit}", F_LLM, color="375623")
    for ci, crit in enumerate(CRITERIA, 5 + len(CRITERIA)):
        hdr_cell(ws, 3, ci, f"{crit}\n{CRIT_NAMES[crit][:4]}", bg_score, color="7B3F00")
    ws.row_dimensions[3].height = 30

    # ── 데이터 행 ────────────────────────────────────────────────────────────
    for rec in samples:
        row = rec["_no"] + 3
        scores = rec["condition_a"]["scores"]

        for ci, val in enumerate([
            rec["_no"], rec.get("grade",""),
            rec.get("subject",""), rec.get("purpose","")
        ], 1):
            data_cell(ws, row, ci, val, F_META,
                      align=center if ci != 4 else left)

        for ci, crit in enumerate(CRITERIA, 5):
            data_cell(ws, row, ci, scores.get(crit,""), F_LLM)

        for ci in range(5 + len(CRITERIA), 5 + 2 * len(CRITERIA)):
            c = ws.cell(row=row, column=ci)
            c.fill = bg_score; c.border = thin
            c.alignment = center

        ws.row_dimensions[row].height = 16

    ws.freeze_panes = "E4"


# ── 결과요약 시트 ─────────────────────────────────────────────────────────────
def make_result_sheet(ws, n_samples):
    ws.sheet_view.showGridLines = False
    for col in range(1, 20):
        ws.column_dimensions[get_column_letter(col)].width = 12

    # ── 제목 ────────────────────────────────────────────────────────────────
    ws.merge_cells("A1:L1")
    c = ws["A1"]
    c.value = "LLM 판사 타당성 검증 결과 요약"
    c.fill = fill("1F3864"); c.font = font(bold=True, size=12, color="FFFFFF")
    c.alignment = center; ws.row_dimensions[1].height = 24

    # ── 섹션 1: 기준별 기술통계 + I-CVI ─────────────────────────────────────
    ws.merge_cells("A3:L3")
    c = ws["A3"]
    c.value = "기준별 평균 적절성 점수 및 내용 타당도 지수 (I-CVI)"
    c.fill = fill("2E4057"); c.font = font(bold=True, size=10, color="FFFFFF")
    c.alignment = center; ws.row_dimensions[3].height = 18

    hdr4 = ["기준", "명칭",
            "R1 평균", "R1 SD",
            "R2 평균", "R2 SD",
            "전체 평균", "전체 SD",
            "R1 I-CVI\n(≥4 비율)", "R2 I-CVI\n(≥4 비율)",
            "I-CVI\n(평균)", "판정\n(≥.78)"]
    for ci, h in enumerate(hdr4, 1):
        hdr_cell(ws, 4, ci, h, fill("404040"), color="FFFFFF")
    ws.row_dimensions[4].height = 30

    # R1 점수: R1_입력 시트, 열 15~24 (5+10 = 15번 컬럼부터)
    # R2 점수: R2_입력 시트, 열 15~24
    r1_cols = [get_column_letter(15 + i) for i in range(len(CRITERIA))]
    r2_cols = [get_column_letter(15 + i) for i in range(len(CRITERIA))]

    for ri, crit in enumerate(CRITERIA):
        row = 5 + ri
        r1_rng = f"R1_입력!{r1_cols[ri]}4:{r1_cols[ri]}{3+n_samples}"
        r2_rng = f"R2_입력!{r2_cols[ri]}4:{r2_cols[ri]}{3+n_samples}"

        vals = [
            crit,
            CRIT_NAMES[crit],
            f"=IFERROR(AVERAGE({r1_rng}),\"\")",
            f"=IFERROR(STDEV({r1_rng}),\"\")",
            f"=IFERROR(AVERAGE({r2_rng}),\"\")",
            f"=IFERROR(STDEV({r2_rng}),\"\")",
            f"=IFERROR((AVERAGE({r1_rng})+AVERAGE({r2_rng}))/2,\"\")",
            f"=IFERROR((STDEV({r1_rng})+STDEV({r2_rng}))/2,\"\")",
            f"=IFERROR(COUNTIF({r1_rng},\">=4\")/{n_samples},\"\")",
            f"=IFERROR(COUNTIF({r2_rng},\">=4\")/{n_samples},\"\")",
            f"=IFERROR((COUNTIF({r1_rng},\">=4\")+COUNTIF({r2_rng},\">=4\"))/(2*{n_samples}),\"\")",
        ]
        bgs = [F_META, F_WHITE, F_R1, F_R1, F_R2, F_R2, F_CALC, F_CALC,
               F_R1, F_R2, F_CALC]
        for ci, (val, bg) in enumerate(zip(vals, bgs), 1):
            c = data_cell(ws, row, ci, formula=val, bg=bg)
            if ci in (3,4,5,6,7,8):
                c.number_format = "0.00"
            if ci in (9,10,11):
                c.number_format = "0.00"

        # 판정 열 (12)
        icvi_cell = f"K{row}"
        c = ws.cell(row=row, column=12,
                    value=f'=IF({icvi_cell}="","",IF({icvi_cell}>=0.78,"O","X"))')
        c.fill = F_CALC; c.font = font(size=9); c.alignment = center; c.border = thin

        ws.row_dimensions[row].height = 16

    # S-CVI 행
    scvi_row = 5 + len(CRITERIA)
    ws.merge_cells(f"A{scvi_row}:B{scvi_row}")
    c = ws.cell(row=scvi_row, column=1, value="S-CVI (전체 평균, 기준 ≥.80)")
    c.fill = fill("1F3864"); c.font = font(bold=True, size=9, color="FFFFFF")
    c.alignment = center; c.border = thin
    ws.cell(row=scvi_row, column=2).border = thin

    scvi_c = ws.cell(row=scvi_row, column=11,
                     value=f"=IFERROR(AVERAGE(K5:K{scvi_row-1}),\"\")")
    scvi_c.fill = fill("BDD7EE"); scvi_c.font = font(bold=True, size=9)
    scvi_c.alignment = center; scvi_c.border = thin; scvi_c.number_format = "0.00"

    judg_c = ws.cell(row=scvi_row, column=12,
                     value=f'=IF(K{scvi_row}="","",IF(K{scvi_row}>=0.8,"O","X"))')
    judg_c.fill = fill("BDD7EE"); judg_c.font = font(bold=True, size=9)
    judg_c.alignment = center; judg_c.border = thin
    ws.row_dimensions[scvi_row].height = 18

    # ── 섹션 2: 두 평가자 상관 ───────────────────────────────────────────────
    sec2_row = scvi_row + 2
    ws.merge_cells(f"A{sec2_row}:L{sec2_row}")
    c = ws[f"A{sec2_row}"]
    c.value = "두 평가자 간 Pearson 상관 (r)  ─  ICC는 Jamovi에서 산출 권장"
    c.fill = fill("2E4057"); c.font = font(bold=True, size=10, color="FFFFFF")
    c.alignment = center; ws.row_dimensions[sec2_row].height = 18

    hdr_r = sec2_row + 1
    for ci, h in enumerate(["기준", "명칭", "Pearson r", "해석"], 1):
        hdr_cell(ws, hdr_r, ci, h, fill("404040"), color="FFFFFF")
    ws.row_dimensions[hdr_r].height = 18

    for ri, crit in enumerate(CRITERIA):
        row = hdr_r + 1 + ri
        r1_rng = f"R1_입력!{r1_cols[ri]}4:{r1_cols[ri]}{3+n_samples}"
        r2_rng = f"R2_입력!{r2_cols[ri]}4:{r2_cols[ri]}{3+n_samples}"
        corr_f = f"=IFERROR(CORREL({r1_rng},{r2_rng}),\"\")"

        data_cell(ws, row, 1, crit, F_META)
        data_cell(ws, row, 2, CRIT_NAMES[crit], F_WHITE, align=left)
        c = data_cell(ws, row, 3, formula=corr_f, bg=F_CALC)
        c.number_format = "0.00"
        interp = ws.cell(row=row, column=4,
            value=f'=IF(C{row}="","",IF(C{row}>=0.9,"매우 높음",IF(C{row}>=0.75,"높음",IF(C{row}>=0.5,"보통","낮음"))))')
        interp.fill = F_CALC; interp.font = font(size=9)
        interp.alignment = center; interp.border = thin
        ws.row_dimensions[row].height = 16

    # 전체 상관
    all_row = hdr_r + 1 + len(CRITERIA)
    r1_all = (f"R1_입력!{r1_cols[0]}4:{r1_cols[-1]}{3+n_samples}")
    r2_all = (f"R2_입력!{r2_cols[0]}4:{r2_cols[-1]}{3+n_samples}")
    ws.merge_cells(f"A{all_row}:B{all_row}")
    c = ws.cell(row=all_row, column=1, value="전체 (10기준 통합)")
    c.fill = fill("1F3864"); c.font = font(bold=True, size=9, color="FFFFFF")
    c.alignment = center; c.border = thin
    ws.cell(row=all_row, column=2).border = thin

    c2 = ws.cell(row=all_row, column=3,
                 value=f"=IFERROR(CORREL({r1_all},{r2_all}),\"\")")
    c2.fill = fill("BDD7EE"); c2.font = font(bold=True, size=9)
    c2.alignment = center; c2.border = thin; c2.number_format = "0.00"

    note = ws.cell(row=all_row + 2, column=1,
        value="※ ICC 산출: Jamovi > Reliability > Inter-rater reliability (ICC) 메뉴 사용\n"
              "   R1·R2 점수 10컬럼을 각각 입력 변수로 지정 (Two-way Mixed, Absolute Agreement)")
    note.font = font(size=8, color="595959"); note.alignment = left
    ws.merge_cells(f"A{all_row+2}:L{all_row+2}")
    ws.row_dimensions[all_row + 2].height = 28


# ── main ─────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    samples = load_samples()
    n = len(samples)
    print(f"샘플 {n}건")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_r1 = wb.create_sheet("R1_입력")
    ws_r2 = wb.create_sheet("R2_입력")
    ws_res = wb.create_sheet("결과요약")

    make_input_sheet(ws_r1, samples, "평가자 1 (R1)", F_R1)
    make_input_sheet(ws_r2, samples, "평가자 2 (R2)", F_R2)
    make_result_sheet(ws_res, n)

    wb.save(OUT_PATH)
    print(f"저장 완료: {OUT_PATH}")
