"""LLM 판사 타당성 검증 — 전문가 검토용 Word 파일 생성.

각 샘플마다:
  - 과제 지시문 / 학생 에세이 / MAFS 피드백
  - 기준별: 루브릭 설명 + LLM 판사 점수 + 판단 근거 + 적절성 5점 척도 입력란

출력: mafs/evaluate/results/validation/expert_review.docx
"""

import json
import random
from pathlib import Path
from collections import defaultdict

RQ1_PATH   = "mafs/evaluate/results/rq1_200.jsonl"
OUT_PATH   = Path("mafs/evaluate/results/validation/expert_review.docx")
SEED       = 42
N_PER_CELL = 1

CRITERIA = ["C1", "C2", "C3", "C4", "C5", "M1", "M2", "M3", "M4", "M5"]

CRIT_INFO = {
    "C1": ("글쓰기 관련성",
           "0점: 에세이 표현을 전혀 인용하지 않거나 1개 이하\n"
           "1점: 에세이 표현 참조는 있으나 따옴표 직접 인용 2개 이하\n"
           "2점: 에세이 표현·문장 5개↑, 그 중 3개↑ 따옴표 직접 인용, 각각 다른 피드백 지점에 연결"),
    "C2": ("약점 구체성",
           "0점: 약점 지적 없거나 '부족하다' 수준의 막연한 서술\n"
           "1점: '위치+이유+인용' 중 일부만 갖추거나, 조건 갖추었으나 1~2건만\n"
           "2점: '위치+이유+에세이 해당 표현 직접 인용' 모두 갖춘 구체적 약점 지적 3건↑"),
    "C3": ("강점 근거",
           "0점: 강점 언급 없거나 '잘 썼다' 수준의 막연한 칭찬\n"
           "1점: 실제 표현 인용 1개이거나, 인용만 있고 '왜 강점인지' 설명 없음\n"
           "2점: 에세이 실제 표현·구조 2개↑ 직접 인용 + 각 인용에 '왜 강점인지' 설명"),
    "C4": ("분석 근거 수치화",
           "0점: 도구 분석 수치(맞춤법 오류 건수, 접속어 유형 수 등)가 전혀 없음\n"
           "1점: 수치가 있으나 학년 기준 없거나 해석 없이 나열만\n"
           "2점: 도구 분석 수치 1개↑ 명시 + 학년 기준과 함께 구체적 제시"),
    "C5": ("학습 성장 전제",
           "0점: 약점 나열만 있고 강점·피드포워드 없음\n"
           "1점: 강점·개선점·피드포워드 중 일부만 있어 균형 무너짐\n"
           "2점: 강점·개선점·피드포워드가 균형 있게 구성, 학습 의욕 충분히 지지"),
    "M1": ("피드포워드",
           "0점: '다음에는 ~해보세요' 형식의 구체적 지침 없음\n"
           "1점: 개선 방향은 있으나 에세이 특정 약점과 연결 안 되거나 형식 모호\n"
           "2점: 에세이의 특정 약점과 연결된 다음 글쓰기 지침이 '다음에는 ~해보세요' 형식으로 명시"),
    "M2": ("분석 타당성",
           "0점: 에세이에 없는 내용 언급 또는 명백한 오분석 1건↑\n"
           "1점: 타당한 분석이 대부분이나 경미한 과장·왜곡 1건\n"
           "2점: 지적한 약점·강점이 에세이 실제 내용에 비추어 모두 타당"),
    "M3": ("학습자 수준 적합",
           "0점: 해당 학년이 이해하기 어려운 어휘·문체·어조\n"
           "1점: 대체로 적합하나 일부 어휘·표현이 학년 수준과 맞지 않음\n"
           "2점: 어휘·문장·어조 전반이 해당 학년 학생이 충분히 이해·수용 가능"),
    "M4": ("과제 맥락 반영",
           "0점: 과제 지시문 조건을 전혀 언급하지 않거나 단순 나열\n"
           "1점: 지시문 조건 1가지만 에세이 내용과 연결하여 분석\n"
           "2점: 과제 지시문 조건을 에세이 내용과 연결하여 2가지↑ 분석 (단순 나열 불인정)"),
    "M5": ("영역별 분석의 충실성",
           "0점: 4개 영역(과제수행·내용·조직·표현) 중 2개↑ 누락\n"
           "1점: 4개 영역이 있으나 일부 영역에 에세이 근거 포함 분석이 1문장 이하\n"
           "2점: 4개 영역 각각에서 에세이 근거를 포함한 분석이 2문장↑"),
}

SCORE_LABEL = {0: "0점 (미충족)", 1: "1점 (부분 충족)", 2: "2점 (완전 충족)"}


# ── 샘플 추출 (make_validation_files.py 와 동일 seed) ─────────────────────────
def load_and_sample():
    records = []
    with open(RQ1_PATH, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                records.append(json.loads(line))

    groups = defaultdict(list)
    for r in records:
        groups[(r.get("grade", ""), r.get("subject", ""))].append(r)

    rng = random.Random(SEED)
    sampled = []
    for key in sorted(groups.keys()):
        sampled.extend(rng.sample(groups[key], min(N_PER_CELL, len(groups[key]))))

    for i, r in enumerate(sampled, 1):
        r["_no"] = i
    return sampled


# ── Word 생성 ─────────────────────────────────────────────────────────────────
def make_word(samples):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def set_bg(cell, hex_color):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(3)

    # ── 표지 ──────────────────────────────────────────────────────────────────
    t = doc.add_paragraph("LLM 판사 채점 적절성 검토")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.runs[0].bold = True
    t.runs[0].font.size = Pt(16)

    sub = doc.add_paragraph("전문가 검토 자료")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(11)

    doc.add_paragraph()

    info_tbl = doc.add_table(rows=3, cols=2)
    info_tbl.style = "Table Grid"
    info_data = [
        ("평가자", ""),
        ("소속 / 전공", ""),
        ("검토일", ""),
    ]
    for ri, (label, val) in enumerate(info_data):
        lc = info_tbl.rows[ri].cells[0]
        vc = info_tbl.rows[ri].cells[1]
        lc.text = label
        vc.text = val
        set_bg(lc, "D9D9D9")
        for p in lc.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)

    doc.add_paragraph()

    guide = doc.add_paragraph(
        "[ 검토 안내 ]\n\n"
        "본 자료는 자동 글쓰기 피드백 시스템(MAFS)의 피드백 품질을 자동 평가한\n"
        "LLM 판사의 채점이 적절한지 검토하기 위해 작성되었습니다.\n\n"
        "각 샘플에 대해 아래 순서로 검토해 주세요:\n"
        "  1) 과제 지시문과 학생 에세이를 읽습니다.\n"
        "  2) MAFS가 생성한 피드백을 읽습니다.\n"
        "  3) 각 기준별로 LLM 판사의 점수와 판단 근거를 확인합니다.\n"
        "  4) LLM 판사의 채점이 적절한지 1~5점으로 평가하고 의견을 기재합니다.\n\n"
        "채점 적절성 척도:\n"
        "  1점 = 전혀 적절하지 않음  /  2점 = 적절하지 않음  /  3점 = 보통\n"
        "  4점 = 적절함  /  5점 = 매우 적절함"
    )
    guide.runs[0].font.size = Pt(9)
    doc.add_page_break()

    # ── 샘플별 내용 ───────────────────────────────────────────────────────────
    for rec in samples:
        no      = rec["_no"]
        grade   = rec.get("grade", "")
        subject = rec.get("subject", "")
        purpose = rec.get("purpose", "")
        prompt  = rec.get("prompt", "").strip()
        essay   = rec.get("essay", "").strip()
        ca      = rec["condition_a"]
        feedback   = ca.get("feedback", "").strip()
        scores     = ca.get("scores", {})
        rationale  = scores.get("rationale", {})

        # 샘플 헤더
        hdr = doc.add_paragraph(
            f"샘플 {no:02d}   |   {grade} · {subject}   |   글쓰기 목적: {purpose}"
        )
        hdr.runs[0].bold = True
        hdr.runs[0].font.size = Pt(12)
        hdr.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        doc.add_paragraph("=" * 60).runs[0].font.size = Pt(8)

        # 과제 지시문
        _block_title(doc, "1. 과제 지시문")
        doc.add_paragraph(prompt).runs[0].font.size = Pt(9)
        doc.add_paragraph()

        # 학생 에세이
        _block_title(doc, "2. 학생 에세이")
        doc.add_paragraph(essay).runs[0].font.size = Pt(9)
        doc.add_paragraph()

        # MAFS 피드백
        _block_title(doc, "3. MAFS 피드백  ← 채점 대상")
        doc.add_paragraph(feedback).runs[0].font.size = Pt(9)
        doc.add_paragraph()

        doc.add_paragraph("─" * 60).runs[0].font.size = Pt(8)
        _block_title(doc, "4. 기준별 LLM 판사 채점 검토")
        doc.add_paragraph(
            "각 기준에 대해 LLM 판사의 점수와 판단 근거를 확인하고,\n"
            "채점이 적절한지 1~5점으로 평가해 주세요."
        ).runs[0].font.size = Pt(9)
        doc.add_paragraph()

        # 기준별 검토 표
        for crit in CRITERIA:
            name, rubric = CRIT_INFO[crit]
            llm_score = scores.get(crit, "-")
            llm_rat   = rationale.get(crit, "(근거 없음)").strip()
            score_label = SCORE_LABEL.get(llm_score, str(llm_score))
            is_trigger = crit in ("C1", "C2", "M1", "M4")
            tag = " ★필수기준" if is_trigger else ""

            # 기준 헤더
            ch = doc.add_paragraph(f"{crit}  {name}{tag}")
            ch.runs[0].bold = True
            ch.runs[0].font.size = Pt(10)
            ch.runs[0].font.color.rgb = (
                RGBColor(0xC0, 0x00, 0x00) if is_trigger
                else RGBColor(0x37, 0x56, 0x23)
            )

            # 기준별 표: 루브릭 / LLM점수 / 근거 / 검토 입력
            tbl = doc.add_table(rows=4, cols=2)
            tbl.style = "Table Grid"

            row_data = [
                ("채점 루브릭",    rubric,      "E8F0FE", "FFFFFF"),
                ("LLM 판사 점수", score_label, "E2EFDA", "E2EFDA"),
                ("LLM 판단 근거", llm_rat,     "FFF9C4", "FFFDE7"),
                ("검토자 평가",
                 "적절성 ( 1 / 2 / 3 / 4 / 5 )점  →  해당 번호에 동그라미 표시\n"
                 "의견:\n\n",
                 "FCE4D6", "FFFFFF"),
            ]
            col_widths = [Cm(3.2), Cm(11.8)]
            for ri, (label, content, lbg, cbg) in enumerate(row_data):
                lc = tbl.rows[ri].cells[0]
                vc = tbl.rows[ri].cells[1]
                lc.text = label
                vc.text = content
                set_bg(lc, lbg)
                set_bg(vc, cbg)
                lc.width = col_widths[0]
                vc.width = col_widths[1]
                for p in lc.paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.size = Pt(9)
                for p in vc.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)

            doc.add_paragraph()

        doc.add_page_break()

    doc.save(OUT_PATH)
    print(f"저장 완료: {OUT_PATH}")


def _block_title(doc, text):
    from docx.shared import Pt, RGBColor
    p = doc.add_paragraph(text)
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)


if __name__ == "__main__":
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    samples = load_and_sample()
    print(f"샘플 {len(samples)}건")
    make_word(samples)
